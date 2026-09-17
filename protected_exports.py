"""Protected register and in-memory password-to-open Office exports.

No source writes, temporary plaintext files, or shared export caches.
"""
from io import BytesIO
from datetime import date, datetime
import hashlib
import re

import pandas as pd


REGISTER_COLUMNS = [
    "Date of Interview/entry", "Child / Beneficiary Name", "Individual number", "Phone number",
    "Specific location", "Gender", "Age group", "Nationality", "Disability",
    "Type of Disability", "Profile status",
]


def _text(value):
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.casefold() in {"nan", "none", "nat", "<na>", "[missing]", "[not recorded]"} else text


def register_gender(value):
    text = _text(value)
    key = text.casefold()
    if key in {"girl", "girls", "woman", "women", "female"}:
        return "Female"
    if key in {"boy", "boys", "man", "men", "male"}:
        return "Male"
    if key == "intersex":
        return "Intersex"
    # In particular, transgender is not a synonym for intersex.
    return text or "Not recorded"


def register_profile(value):
    text = _text(value)
    key = text.casefold().replace("_", " ")
    if key == "host" or key.startswith("host community"):
        return "Host"
    if key.startswith("refugee"):
        return "Refugee"
    return text or "Not recorded"


def beneficiary_register(frame, age_mapper):
    """All supplied filtered entries, not a deduplicated beneficiary roster."""
    if frame.empty:
        return pd.DataFrame(columns=REGISTER_COLUMNS)
    source = frame.reset_index(drop=True)

    def column(name):
        return source.get(name, pd.Series(pd.NA, index=source.index))

    def dates(name):
        return (pd.to_datetime(column(name), errors="coerce", format="mixed", utc=True)
                .dt.tz_convert("Africa/Nairobi").dt.tz_localize(None).dt.normalize())

    result = pd.DataFrame(index=source.index)
    result[REGISTER_COLUMNS[0]] = dates("interview_date").combine_first(dates("reporting_date"))
    result[REGISTER_COLUMNS[1]] = column("information_seeker_name").map(_text).replace("", "Not recorded")
    result[REGISTER_COLUMNS[2]] = column("information_seeker_individual_number").map(_text).replace("", "Not recorded")
    primary_phone = column("information_seeker_phone").map(_text)
    alternative_phone = column("alternative_phone").map(_text)
    missing_phone_labels = {"not recorded", "not provided", "n/a", "na"}
    primary_phone = primary_phone.mask(primary_phone.str.casefold().isin(missing_phone_labels), "")
    alternative_phone = alternative_phone.mask(alternative_phone.str.casefold().isin(missing_phone_labels), "")
    result["Phone number"] = primary_phone.mask(primary_phone.eq(""), alternative_phone).replace("", "Not recorded")
    # Keep both source contexts explicit: helpdesk Section/Block is not residence.
    section = column("helpdesk_section_block").map(_text)
    residence = column("residence_neighborhood_compound_house").map(_text)
    result["Specific location"] = [
        "; ".join(part for part in (
            f"Section/Block: {block}" if block else "",
            f"Neighborhood/Compound/House: {home}" if home else "",
        ) if part) or "Not recorded"
        for block, home in zip(section, residence)
    ]
    result["Gender"] = column("information_seeker_gender").map(register_gender)
    age = column("age_group").combine_first(column("information_seeker_age"))
    result["Age group"] = age.map(age_mapper).map(_text).str.replace(r" years$", " Yrs", regex=True).replace("", "Not recorded")
    nationality = column("information_seeker_nationality").map(_text)
    other = column("information_seeker_nationality_other").map(_text)
    use_other = nationality.str.casefold().isin(["", "other", "others", "other nationality", "other not listed", "not listed"])
    result["Nationality"] = nationality.mask(use_other & other.ne(""), other).replace("", "Not recorded")
    result["Disability"] = column("disability_status").map(_text).replace("", "Not recorded")
    result["Type of Disability"] = column("disability_type").map(_text).replace("", "Not recorded")
    result["Profile status"] = column("household_type").map(register_profile)
    return result[REGISTER_COLUMNS].sort_values(REGISTER_COLUMNS[0], ascending=False, na_position="last", kind="stable").reset_index(drop=True)


def encrypted_excel_bytes(table, password):
    """Return only verified Office Agile encrypted XLSX bytes, or raise.

    Sheet/workbook protection is not encryption. The OOXML file itself is
    encrypted; a password is required before its data can be opened.
    """
    if not isinstance(password, str) or len(password) < 12 or not password.strip():
        raise ValueError("Use a file password with at least 12 characters.")
    if len(table) > 1_048_575 or len(table.columns) > 16_384:
        raise ValueError("Selection exceeds Excel limits. Narrow the report filters.")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    import msoffcrypto
    from msoffcrypto.format.ooxml import OOXMLFile

    workbook = Workbook()
    workbook.properties.creator = "TDH Kenya"
    sheet = workbook.active
    sheet.title = "Protected records"
    sheet.freeze_panes = "A2"

    def write_cell(row, col, value):
        cell = sheet.cell(row, col)
        if value is None or pd.isna(value):
            return cell
        if isinstance(value, (datetime, date, pd.Timestamp)):
            stamp = pd.Timestamp(value)
            if stamp.tzinfo is not None:
                stamp = stamp.tz_convert("Africa/Nairobi").tz_localize(None)
            cell.value = stamp.to_pydatetime()
            cell.number_format = "dd mmm yyyy"
        elif isinstance(value, str):
            if len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
                raise ValueError("A field cannot be represented safely in Excel.")
            cell.value = value
            cell.data_type = "s"  # Never execute source text beginning with =, +, -, @.
            cell.number_format = "@"
        else:
            cell.value = value
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        return cell

    for col, name in enumerate(table.columns, 1):
        cell = write_cell(1, col, str(name))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F7D69")
        sheet.column_dimensions[get_column_letter(col)].width = 30 if "name" in str(name).casefold() or "location" in str(name).casefold() else 23
    sheet.row_dimensions[1].height = 32
    for row, values in enumerate(table.itertuples(index=False, name=None), 2):
        for col, value in enumerate(values, 1):
            write_cell(row, col, value)
    sheet.auto_filter.ref = sheet.dimensions
    plain, encrypted, checked = BytesIO(), BytesIO(), BytesIO()
    try:
        workbook.save(plain)
        plain.seek(0)
        OOXMLFile(plain).encrypt(password, encrypted)
        encrypted.seek(0)
        verifier = msoffcrypto.OfficeFile(encrypted)
        if not verifier.is_encrypted():
            raise RuntimeError("Export encryption verification failed.")
        verifier.load_key(password=password, verify_password=True)
        verifier.decrypt(checked, verify_integrity=True)
        if hashlib.sha256(checked.getvalue()).digest() != hashlib.sha256(plain.getvalue()).digest():
            raise RuntimeError("Export integrity verification failed.")
        return encrypted.getvalue()
    finally:
        workbook.close()
        plain.close()
        encrypted.close()
        checked.close()
